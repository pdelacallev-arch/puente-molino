import argparse
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = str(PROJECT_ROOT / "plantillas" / "plantilla_RFIs.docx")
MARKDOWN = str(PROJECT_ROOT / "entregables" / "04_borradores" / "RFI-EST-2026-002-R00.md")
OUTPUT = str(PROJECT_ROOT / "entregables" / "04_borradores" / "RFI-EST-2026-002-R00.docx")

FONT = "Calibri"
BLACK = RGBColor(0x00, 0x00, 0x00)
HEADER_BG = "D9D9D9"
BORDER_COLOR = "BFBFBF"
BODY_SIZE = Pt(9)
TABLE_SIZE = Pt(7.5)
CALLOUT_SIZE = Pt(8.5)
HEADING_2_SIZE = Pt(12)
HEADING_3_SIZE = Pt(10.5)
RESPONSE_BODY_SIZE = Pt(10.5)
RESPONSE_TABLE_SIZE = Pt(9)
RESPONSE_HEADING_3_SIZE = Pt(11.5)


def parse_md(filepath):
    with open(filepath, encoding="utf-8") as f:
        return f.readlines()


def is_table_line(line):
    return line.strip().startswith("|") and line.strip().endswith("|")


def parse_table(lines, i):
    while i < len(lines) and not lines[i].strip().startswith("|"):
        i += 1
    if i >= len(lines):
        return None, i
    headers = [h.strip() for h in lines[i].strip().split("|")[1:-1]]
    i += 1
    if i < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i].strip()):
        i += 1
    rows = []
    while i < len(lines):
        s = lines[i].strip()
        if not s.startswith("|"):
            break
        cells = [c.strip() for c in s.split("|")[1:-1]]
        if cells:
            rows.append(cells)
        i += 1
    return {"headers": headers, "rows": rows}, i


def shade(cell, color):
    cell._tc.get_or_add_tcPr().append(
        parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    )


def cell_margin(cell, top=20, bottom=20, left=40, right=40):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:start w:w="{left}" w:type="dxa"/>'
        f'  <w:end w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    ))


def table_borders(tbl):
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'</w:tblBorders>'
    ))


def keep_next(p):
    p._element.get_or_add_pPr().append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))


def keep_lines(p):
    p._element.get_or_add_pPr().append(parse_xml(f'<w:keepLines {nsdecls("w")}/>'))


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def set_cell_vertical_center(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))


def autofit_table(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="0" w:type="auto"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="autofit"/>'))


def set_table_geometry(tbl, widths_cm):
    """Apply exact widths so Word does not redistribute narrow RFI columns."""
    widths_dxa = [round(width * 567) for width in widths_cm]
    total_dxa = sum(widths_dxa)
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    for tag in ("tblW", "tblLayout", "tblInd"):
        for node in list(tblPr.findall(qn(f"w:{tag}"))):
            tblPr.remove(node)
    tblPr.append(parse_xml(
        f'<w:tblW {nsdecls("w")} w:w="{total_dxa}" w:type="dxa"/>'
    ))
    tblPr.append(parse_xml(
        f'<w:tblInd {nsdecls("w")} w:w="0" w:type="dxa"/>'
    ))
    tblPr.append(parse_xml(
        f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'
    ))

    grid = tbl._tbl.tblGrid
    for node in list(grid):
        grid.remove(node)
    for width in widths_dxa:
        grid.append(parse_xml(
            f'<w:gridCol {nsdecls("w")} w:w="{width}"/>'
        ))

    for row in tbl.rows:
        for cell, width in zip(row.cells, widths_dxa):
            tcPr = cell._tc.get_or_add_tcPr()
            for node in list(tcPr.findall(qn("w:tcW"))):
                tcPr.remove(node)
            tcPr.insert(0, parse_xml(
                f'<w:tcW {nsdecls("w")} w:w="{width}" w:type="dxa"/>'
            ))


def table_widths(headers):
    """Return deliberate widths for the recurring RFI table families."""
    if headers == ["Campo", "Información", "Campo", "Información"]:
        return [2.6, 7.0, 2.8, 3.6]
    if len(headers) == 4 and any("Definición requerida" in h for h in headers):
        return [2.7, 4.1, 5.3, 3.9]
    if len(headers) == 6:
        return [2.7, 2.0, 2.0, 2.2, 3.8, 3.3]
    if len(headers) == 2:
        return [7.0, 9.0]
    if len(headers) == 3:
        return [5.33, 5.34, 5.33]
    return [16.0 / len(headers)] * len(headers)


def add_math_runs(paragraph, text, size, bold=False, italic=False):
    """Render simple inline LaTeX variables without exposing Markdown syntax."""
    text = text.replace(r"\,", " ").replace("\\", "")
    cursor = 0
    pattern = re.compile(r"([A-Za-z])_\{?([A-Za-z]+)\}?")
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            run.bold = bold
            run.italic = italic
            run.font.size = size
            run.font.name = FONT
            run.font.color.rgb = BLACK
        base = paragraph.add_run(match.group(1))
        base.bold = bold
        base.italic = True
        base.font.size = size
        base.font.name = FONT
        base.font.color.rgb = BLACK
        sub = paragraph.add_run(match.group(2))
        sub.bold = bold
        sub.italic = True
        sub.font.subscript = True
        sub.font.size = size
        sub.font.name = FONT
        sub.font.color.rgb = BLACK
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.bold = bold
        run.italic = italic
        run.font.size = size
        run.font.name = FONT
        run.font.color.rgb = BLACK


def add_inline_runs(paragraph, text, size, bold=False, italic=False):
    """Render the inline Markdown used by project RFIs."""
    text = text.replace(r'\"', '"')
    parts = re.split(r'(\*\*.*?\*\*|\*[^*]+?\*|`[^`]+`|\$[^$]+?\$)', text)
    for part in parts:
        if not part:
            continue
        part_bold = bold
        part_italic = italic
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            part_bold = True
        elif part.startswith("*") and part.endswith("*"):
            part = part[1:-1]
            part_italic = True
        elif part.startswith("`") and part.endswith("`"):
            part = part[1:-1]
        elif part.startswith("$") and part.endswith("$"):
            add_math_runs(paragraph, part[1:-1], size, part_bold, True)
            continue
        run = paragraph.add_run(part)
        run.bold = part_bold
        run.italic = part_italic
        run.font.size = size
        run.font.name = FONT
        run.font.color.rgb = BLACK


def merge_cells(t, start_row, start_col, end_row, end_col):
    """Merge cells in a rectangular region."""
    for ri in range(start_row, end_row + 1):
        for ci in range(start_col, end_col + 1):
            cell = t.rows[ri].cells[ci]
            if ri == start_row and ci == start_col:
                continue  # skip the anchor cell
            cell._tc.getparent().remove(cell._tc)


def render(markdown_path=None, output_path=None, template_path=None):
    markdown_path = Path(markdown_path or MARKDOWN)
    output_path = Path(output_path or OUTPUT)
    template_path = Path(template_path or TEMPLATE)
    lines = parse_md(markdown_path)
    doc = Document(template_path)

    # ── margins 2.5 cm ──
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)

    # ── configure styles ──
    def cfg(name, fname=None, size=None, bold=None, color=None,
            space_before=None, space_after=None, line_spacing=None,
            alignment=None):
        s = doc.styles[name]
        if fname:
            s.font.name = fname
            rPr = s.element.get_or_add_rPr()
            rf = rPr.find(qn('w:rFonts'))
            if rf is None:
                rf = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
                rPr.insert(0, rf)
            rf.set(qn('w:eastAsia'), fname)
        if size:
            s.font.size = size
        if bold is not None:
            s.font.bold = bold
        if color:
            s.font.color.rgb = color
        if space_before is not None:
            s.paragraph_format.space_before = space_before
        if space_after is not None:
            s.paragraph_format.space_after = space_after
        if line_spacing is not None:
            s.paragraph_format.line_spacing = line_spacing
        if alignment is not None:
            s.paragraph_format.alignment = alignment

    cfg("Normal", fname=FONT, size=BODY_SIZE, color=BLACK,
        line_spacing=1.0, space_after=Pt(3), space_before=Pt(0))
    cfg("Heading 1", fname=FONT, size=Pt(13), bold=True, color=BLACK,
        space_before=Pt(8), space_after=Pt(4), alignment=WD_ALIGN_PARAGRAPH.LEFT)
    cfg("Heading 2", fname=FONT, size=HEADING_2_SIZE, bold=True, color=BLACK,
        space_before=Pt(6), space_after=Pt(4), alignment=WD_ALIGN_PARAGRAPH.CENTER)
    cfg("Heading 3", fname=FONT, size=HEADING_3_SIZE, bold=True, color=BLACK,
        space_before=Pt(6), space_after=Pt(3), alignment=WD_ALIGN_PARAGRAPH.LEFT)

    # The source headings already contain their documentary numbering. Remove
    # template-linked numbering to prevent output such as "1. 1. Motivo".
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        pPr = doc.styles[style_name].element.get_or_add_pPr()
        numPr = pPr.find(qn("w:numPr"))
        if numPr is not None:
            pPr.remove(numPr)

    # Configure RFI Heading 1 style if it exists
    if "RFI Heading 1" in [s.name for s in doc.styles]:
        cfg("RFI Heading 1", fname=FONT, size=Pt(11), bold=True, color=BLACK,
            space_before=Pt(6), space_after=Pt(4), alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # ── numbering ──
    numbering_part = doc.part.numbering_part
    nlm = numbering_part.element
    for child in list(nlm):
        tag = child.tag
        if 'abstractNum' in tag or 'num' in tag:
            nlm.remove(child)

    abs_num = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="0">'
        f'  <w:multiLevelType w:val="singleLevel"/>'
        f'  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/>'
        f'    <w:lvlText w:val="•"/><w:lvlJc w:val="left"/>'
        f'    <w:pPr><w:tabs><w:tab w:val="num" w:pos="360"/></w:tabs><w:ind w:left="360" w:hanging="180"/></w:pPr>'
        f'  </w:lvl>'
        f'</w:abstractNum>'
    )
    nlm.append(abs_num)
    num_elm = parse_xml(
        f'<w:num {nsdecls("w")} w:numId="1">'
        f'  <w:abstractNumId w:val="0"/>'
        f'</w:num>'
    )
    nlm.append(num_elm)

    # ── clean body ──
    body = doc.element.body
    for child in list(body):
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)

    # ── write content ──
    i = 0
    in_response = False
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()

        if not s:
            i += 1
            continue

        # ── headings ──
        if s.startswith("#"):
            level = len(s.split(" ")[0])
            text = s[level:].strip()

            # H1: ANEXO 1 — RFI-EST-2026-001-R00
            # H2: FICHA DE SOLICITUD DE INFORMACIÓN, RESPUESTA DEL PROYECTISTA
            # H3: RFI-EST-2026-001 · REVISIÓN R00, 1. Motivo, 2. Condición, etc.
            if level == 1:
                style = "Heading 1"
            elif level == 2:
                style = "Heading 2"
            else:
                # H3 — check if it's "RFI-EST-2026-001 · REVISIÓN R00" (no number)
                # or numbered sections like "1. Motivo de la consulta"
                style = "Heading 3"

            p = doc.add_paragraph(style=style)

            if level == 2 or (level == 3 and text.startswith("RFI-")):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if in_response and level == 3:
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(5)

            run = p.add_run(text)
            sizes = {1: Pt(13), 2: HEADING_2_SIZE, 3: HEADING_3_SIZE}
            if in_response and level == 3:
                sizes[3] = RESPONSE_HEADING_3_SIZE
            run.bold = True
            run.font.size = sizes.get(level, HEADING_3_SIZE)
            run.font.name = FONT
            run.font.color.rgb = BLACK

            keep_next(p)
            i += 1
            continue

        # ── tables ──
        if is_table_line(s):
            result, nxt = parse_table(lines, i)
            if result:
                hdrs = result["headers"]
                rows = result["rows"]
                nc = len(hdrs)
                nr = 1 + len(rows)

                t = doc.add_table(rows=nr, cols=nc)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                table_borders(t)
                set_table_geometry(t, table_widths(hdrs))
                table_size = RESPONSE_TABLE_SIZE if in_response else TABLE_SIZE
                margin_args = (30, 30, 50, 50) if in_response else (20, 20, 40, 40)

                # header row
                row_header(t.rows[0])
                for ci, h in enumerate(hdrs):
                    cell = t.rows[0].cells[ci]
                    cell.text = ""
                    set_cell_vertical_center(cell)
                    cell_margin(cell, *margin_args)
                    shade(cell, HEADER_BG)

                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(1)
                    p.paragraph_format.space_after = Pt(1)
                    add_inline_runs(p, h, table_size, bold=True)

                # data rows
                for ri, rd in enumerate(rows):
                    try:
                        row_cant_split(t.rows[ri + 1])
                    except Exception:
                        pass
                    for ci, val in enumerate(rd):
                        if ci >= nc:
                            continue
                        cell = t.rows[ri + 1].cells[ci]
                        cell.text = ""
                        set_cell_vertical_center(cell)
                        cell_margin(cell, *margin_args)

                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(1)
                        p.paragraph_format.space_after = Pt(1)
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        # Detect checkbox columns: center them
                        is_checkbox = val.strip() in ["☐", "", "☐ Sí · ☐ No · ☐ Condicionado. Alcance:"]
                        if is_checkbox or re.match(r'^☐', val.strip()):
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                        # Bold for ID cells and special rows
                        force_bold = (
                            ci == 0
                            or val.strip().startswith("CONSULTA")
                            or val.strip().startswith("RESTRICCIÓN")
                            or val.strip().startswith("EMISIÓN")
                            or val.strip() == "RECEPCIÓN DE LA RESPUESTA"
                        )
                        add_inline_runs(p, val, table_size, bold=force_bold)

                # Compact response forms should move as a unit instead of
                # leaving two rows on one page and the remainder on the next.
                if hdrs == ["Definición", "Respuesta"]:
                    for row in t.rows[:-1]:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                keep_next(paragraph)

                # spacer
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after = Pt(2)
                r = sp.add_run("")
                r.font.size = Pt(1)

                i = nxt
                continue
            else:
                i += 1
                continue

        # ── blockquotes ──
        if s.startswith(">"):
            text = s[1:].strip()
            p = doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline_runs(p, text, CALLOUT_SIZE, italic=True)
            pPr = p._element.get_or_add_pPr()
            pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="567" w:right="567"/>'))
            pPr.append(parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:left w:val="single" w:sz="8" w:space="6" w:color="C0C0C0"/>'
                f'</w:pBdr>'
            ))
            keep_lines(p)
            i += 1
            continue

        # ── bullet lists ──
        if re.match(r"^[-*]\s+", s):
            text = re.sub(r"^[-*]\s+", "", s)
            p = doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pPr = p._element.get_or_add_pPr()
            pPr.append(parse_xml(
                f'<w:numPr {nsdecls("w")}>'
                f'  <w:ilvl w:val="0"/>'
                f'  <w:numId w:val="1"/>'
                f'</w:numPr>'
            ))
            add_inline_runs(p, text, BODY_SIZE)
            i += 1
            continue

        # ── horizontal rule ──
        if s.startswith("---") and len(s) >= 3:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_break(WD_BREAK.PAGE)
            in_response = True
            i += 1
            continue

        # ── regular paragraph ──
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if in_response:
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.space_after = Pt(5)

        paragraph_size = RESPONSE_BODY_SIZE if in_response else BODY_SIZE
        add_inline_runs(p, s, paragraph_size)

        i += 1

    # ── footer page numbers ──
    for sec in doc.sections:
        footer = sec.footer
        footer.is_linked_to_previous = False
        for fp in footer.paragraphs:
            fp.clear()

        if footer.paragraphs:
            fp = footer.paragraphs[0]
        else:
            fp = footer.add_paragraph()
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(4)

        def fld(code):
            r = fp.add_run()
            r._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
            r2 = fp.add_run()
            r2._element.append(parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> {code} </w:instrText>'))
            r3 = fp.add_run()
            r3._element.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))

        fp.add_run("Página ")
        fld("PAGE")
        fp.add_run(" de ")
        fld("NUMPAGES")

        for r in fp.runs:
            r.font.size = Pt(7.5)
            r.font.name = FONT
            r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Documento generado: {output_path.resolve()}")
    return output_path


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Renderiza un RFI Markdown como DOCX de dos paginas."
    )
    parser.add_argument("--input", required=True, help="Archivo RFI Markdown de entrada.")
    parser.add_argument(
        "--output",
        help="DOCX de salida. Por defecto usa el mismo nombre del Markdown.",
    )
    parser.add_argument(
        "--template",
        default=TEMPLATE,
        help="Plantilla DOCX; por defecto plantillas/plantilla_RFIs.docx.",
    )
    args = parser.parse_args(argv)
    input_path = Path(args.input)
    output_path = Path(args.output) if args.output else input_path.with_suffix(".docx")
    render(input_path, output_path, args.template)


if __name__ == "__main__":
    main()
