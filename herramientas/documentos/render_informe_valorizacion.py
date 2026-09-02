#!/usr/bin/env python3
"""
render_informe_valorizacion.py - Renderiza el Informe de Valorización / Actividades
desde Markdown a DOCX según la plantilla IES2 y confi_aspecto.txt.

Cumple con:
- Plantilla: plantillas/plantilla_IES2.docx
- Aspecto: Calibri 11 pt, 1.15 interlineado, 0 pt antes / 6 pt después, justificado
- Títulos: H1 16pt, H2 14pt, H3 12pt bold, con keepNext
- Tablas: Encabezado #D9D9D9 bold, bordes finos #BFBFBF, cantSplit, tblHeader, autofit
- Imágenes: Centradas con resolución y ancho adecuado, pie de foto keepWithNext
"""

import os
import re
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

PROJECT_ROOT = Path(__file__).resolve().parents[2]
BASE = str(PROJECT_ROOT)
TEMPLATE = os.path.join(BASE, "plantillas", "plantilla_IES2.docx")
DEFAULT_MARKDOWN = os.path.join(BASE, "informes_actividades", "informe_valorizacion_agosto_2026.md")
DEFAULT_OUTPUT = os.path.join(BASE, "informes_actividades", "informe_valorizacion_agosto_2026.docx")

FONT = "Calibri"
BLACK = RGBColor(0x00, 0x00, 0x00)
HEADER_BG = "D9D9D9"
BORDER_COLOR = "BFBFBF"
IMG_WIDTH_CM = 14.0

IMG_RE = re.compile(r'!\[([^\]]*)\]\((?:<([^>]+)>|([^)]+))\)')
TABLE_LINE_RE = re.compile(r'^\|(.+)\|$')
SEPARATOR_RE = re.compile(r'^\|[\s\-:|]+\|$')


def parse_md(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        return f.readlines()


def resolve_image_path(img_relpath, md_dir):
    if os.path.isabs(img_relpath):
        return img_relpath
    clean_path = img_relpath.replace('/', os.sep).replace('\\', os.sep)
    return os.path.normpath(os.path.join(md_dir, clean_path))


def is_table_line(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|")


def parse_table(lines, i):
    while i < len(lines) and not is_table_line(lines[i]):
        i += 1
    if i >= len(lines):
        return None, i
    
    header_line = lines[i].strip()
    headers = [c.strip() for c in header_line[1:-1].split("|")]
    i += 1
    
    if i < len(lines) and SEPARATOR_RE.match(lines[i].strip()):
        i += 1
        
    rows = []
    while i < len(lines):
        s = lines[i].strip()
        if not is_table_line(s):
            break
        cells = [c.strip() for c in s[1:-1].split("|")]
        if cells:
            rows.append(cells)
        i += 1
    return {"headers": headers, "rows": rows}, i


# ── XML / Estilo Helpers ──

def shade(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))


def cell_margin(cell, top=40, bottom=40, left=70, right=70):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:start w:w="{left}" w:type="dxa"/>'
        f'  <w:end w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    ))


def table_borders(tbl):
    tbl._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'</w:tblBorders>'
    ))


def keep_next(p):
    p._element.get_or_add_pPr().append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))


def keep_lines(p):
    p._element.get_or_add_pPr().append(parse_xml(f'<w:keepLines {nsdecls("w")}/>'))


def row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))


def row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))


def set_cell_vertical_center(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))


def page_break_before(p):
    pPr = p._element.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pageBreakBefore {nsdecls("w")}/>'))


def autofit_table(tbl):
    tblPr = tbl._tbl.tblPr
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="0" w:type="auto"/>'))
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="autofit"/>'))


def add_formatted_runs(paragraph, text, base_size=Pt(11), base_color=BLACK, italic=False):
    """Parsea markdown simple (**negrita**, *cursiva*, `código`, <br>) y agrega runs."""
    if not text:
        return
    
    # Manejar <br>
    lines = re.split(r'<br\s*/?>', text, flags=re.IGNORECASE)
    for l_idx, line_text in enumerate(lines):
        if l_idx > 0:
            paragraph.add_run('\n')
        
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', line_text)
        for token in tokens:
            if not token:
                continue
            if token.startswith("**") and token.endswith("**") and len(token) >= 4:
                run = paragraph.add_run(token[2:-2])
                run.bold = True
                run.italic = italic
            elif token.startswith("*") and token.endswith("*") and len(token) >= 2:
                run = paragraph.add_run(token[1:-1])
                run.italic = True
            elif token.startswith("`") and token.endswith("`") and len(token) >= 2:
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
                run.italic = italic
            else:
                # Limpiar enlaces markdown [texto](url) si los hubiera
                clean = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', token)
                run = paragraph.add_run(clean)
                run.italic = italic
            
            run.font.name = FONT
            run.font.size = base_size
            run.font.color.rgb = base_color


def render_document(md_path=DEFAULT_MARKDOWN, out_path=DEFAULT_OUTPUT, template_path=TEMPLATE):
    lines = parse_md(md_path)
    md_dir = os.path.dirname(os.path.abspath(md_path))
    
    if os.path.exists(template_path):
        doc = Document(template_path)
    else:
        doc = Document()

    # ── Configurar Márgenes A4 (2.5 cm) ──
    for sec in doc.sections:
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)

    # ── Configurar Estilos ──
    def cfg_style(name, fname=FONT, size=Pt(11), bold=False, color=BLACK,
                  space_before=Pt(0), space_after=Pt(6), line_spacing=1.15,
                  alignment=WD_ALIGN_PARAGRAPH.LEFT):
        if name in doc.styles:
            s = doc.styles[name]
        else:
            s = doc.styles.add_style(name, 1)
        s.font.name = fname
        s.font.size = size
        s.font.bold = bold
        s.font.color.rgb = color
        s.paragraph_format.space_before = space_before
        s.paragraph_format.space_after = space_after
        s.paragraph_format.line_spacing = line_spacing
        s.paragraph_format.alignment = alignment

    cfg_style("Normal", size=Pt(11), line_spacing=1.15, space_after=Pt(6), alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    cfg_style("Heading 1", size=Pt(16), bold=True, space_before=Pt(20), space_after=Pt(8), alignment=WD_ALIGN_PARAGRAPH.LEFT)
    cfg_style("Heading 2", size=Pt(14), bold=True, space_before=Pt(14), space_after=Pt(6), alignment=WD_ALIGN_PARAGRAPH.LEFT)
    cfg_style("Heading 3", size=Pt(12), bold=True, space_before=Pt(10), space_after=Pt(4), alignment=WD_ALIGN_PARAGRAPH.LEFT)

    # ── Limpiar cuerpo inicial de la plantilla ──
    body = doc.element.body
    for child in list(body):
        if child.tag in (qn('w:p'), qn('w:tbl')):
            body.remove(child)

    i = 0
    h1_count = 0
    
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()

        if not s:
            i += 1
            continue

        # ── Separador / Línea horizontal ──
        if s.startswith("___") or s.startswith("---"):
            p = doc.add_paragraph()
            pPr = p._element.get_or_add_pPr()
            pPr.append(parse_xml(
                f'<w:pBdr {nsdecls("w")}>'
                f'  <w:bottom w:val="single" w:sz="6" w:space="4" w:color="BBBBBB"/>'
                f'</w:pBdr>'
            ))
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(8)
            i += 1
            continue

        # ── Encabezados (# H1, ## H2, ### H3) ──
        if s.startswith("#"):
            match = re.match(r'^(#+)\s*(.*)$', s)
            if match:
                hashes, title_text = match.groups()
                level = len(hashes)
                style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}.get(level, "Heading 3")
                
                p = doc.add_paragraph(style=style)
                
                # Ajustar salto antes de H1 principal (si es nuevo capítulo)
                if level == 1 and h1_count > 0:
                    page_break_before(p)
                if level == 1:
                    h1_count += 1
                
                add_formatted_runs(p, title_text, base_size={1: Pt(16), 2: Pt(14), 3: Pt(12)}.get(level, Pt(12)), base_color=BLACK)
                p.runs[0].bold = True if p.runs else None
                keep_next(p)
                i += 1
                continue

        # ── Imágenes en Markdown ──
        img_match = IMG_RE.match(s)
        if img_match:
            alt_text = img_match.group(1)
            img_relpath = img_match.group(2) or img_match.group(3)
            img_path = resolve_image_path(img_relpath, md_dir)
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            
            if os.path.exists(img_path):
                run = p.add_run()
                run.add_picture(img_path, width=Cm(IMG_WIDTH_CM))
            else:
                run = p.add_run(f"[Imagen no encontrada: {alt_text}]")
                run.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
            
            keep_next(p)
            i += 1
            continue

        # ── Título de Tabla / Pie de Figura ──
        if s.startswith("**Tabla N.º") or s.startswith("**Tabla N°") or s.startswith("**Fotografía N.º"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            if "Tabla" in s:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                keep_next(p)
                add_formatted_runs(p, s, base_size=Pt(10), base_color=BLACK)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                add_formatted_runs(p, s, base_size=Pt(9.5), base_color=BLACK)
            i += 1
            continue

        # ── Tablas ──
        if is_table_line(s):
            table_data, next_idx = parse_table(lines, i)
            if table_data:
                headers = table_data["headers"]
                rows = table_data["rows"]
                num_cols = len(headers)
                num_rows = 1 + len(rows)

                tbl = doc.add_table(rows=num_rows, cols=num_cols)
                tbl.style = "Table Grid"
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                table_borders(tbl)
                autofit_table(tbl)

                # Celda header
                header_row = tbl.rows[0]
                row_header(header_row)
                row_cant_split(header_row)

                # Detectar tamaño de fuente según cantidad de columnas
                tbl_font_size = Pt(8) if num_cols >= 5 else Pt(8.5)

                for c_idx, h_text in enumerate(headers):
                    cell = header_row.cells[c_idx]
                    cell.text = ""
                    set_cell_vertical_center(cell)
                    cell_margin(cell, top=50, bottom=50, left=60, right=60)
                    shade(cell, HEADER_BG)

                    hp = cell.paragraphs[0]
                    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    hp.paragraph_format.space_before = Pt(2)
                    hp.paragraph_format.space_after = Pt(2)
                    add_formatted_runs(hp, h_text, base_size=tbl_font_size, base_color=BLACK)
                    for r in hp.runs:
                        r.bold = True

                # Celdas datos
                for r_idx, row_cells in enumerate(rows):
                    table_row = tbl.rows[r_idx + 1]
                    row_cant_split(table_row)

                    for c_idx, cell_value in enumerate(row_cells):
                        if c_idx >= num_cols:
                            continue
                        cell = table_row.cells[c_idx]
                        cell.text = ""
                        set_cell_vertical_center(cell)
                        cell_margin(cell, top=40, bottom=40, left=60, right=60)

                        dp = cell.paragraphs[0]
                        dp.paragraph_format.space_before = Pt(1)
                        dp.paragraph_format.space_after = Pt(1)
                        
                        # Alinear al centro si es ID, N.º, fecha corta
                        if c_idx == 0 or (len(cell_value) <= 12 and re.match(r'^\d{2}/\d{2}/\d{4}$|^RP-\d+|^01\.', cell_value.strip())):
                            dp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        else:
                            dp.alignment = WD_ALIGN_PARAGRAPH.LEFT

                        add_formatted_runs(dp, cell_value, base_size=tbl_font_size, base_color=BLACK)

                # Espacio tras tabla
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(2)
                sp.paragraph_format.space_after = Pt(4)
                
                i = next_idx
                continue
            else:
                i += 1
                continue

        # ── Listas con viñetas o numeradas ──
        bullet_match = re.match(r'^(\-|\*)\s+(.*)$', s)
        numbered_match = re.match(r'^(\d+)\.\s+(.*)$', s)
        
        if bullet_match:
            item_text = bullet_match.group(2)
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            run_bullet = p.add_run("•  ")
            run_bullet.bold = True
            run_bullet.font.name = FONT
            run_bullet.font.size = Pt(11)
            add_formatted_runs(p, item_text, base_size=Pt(11), base_color=BLACK)
            i += 1
            continue
            
        elif numbered_match and not s.startswith("##"):
            num_str = numbered_match.group(1)
            item_text = numbered_match.group(2)
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4)
            run_num = p.add_run(f"{num_str}. ")
            run_num.bold = True
            run_num.font.name = FONT
            run_num.font.size = Pt(11)
            add_formatted_runs(p, item_text, base_size=Pt(11), base_color=BLACK)
            i += 1
            continue

        # ── Párrafo normal ──
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        add_formatted_runs(p, s, base_size=Pt(11), base_color=BLACK)
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    print(f"Documento DOCX generado exitosamente en:\n{out_path}")
    return out_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Renderiza informe de actividades / valorización a DOCX.")
    parser.add_argument("--input", default=DEFAULT_MARKDOWN, help="Ruta al archivo Markdown de entrada")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, help="Ruta al archivo DOCX de salida")
    parser.add_argument("--template", default=TEMPLATE, help="Ruta a la plantilla DOCX")
    args = parser.parse_args()

    render_document(args.input, args.output, args.template)
