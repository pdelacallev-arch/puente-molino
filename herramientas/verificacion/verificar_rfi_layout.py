import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


MAIN_TITLE = "FICHA DE SOLICITUD DE INFORMACIÓN"
RESPONSE_TITLE = "PRONUNCIAMIENTO DE LA SUPERVISIÓN"


def validate(docx_path):
    doc = Document(docx_path)
    errors = []
    paragraphs = {p.text.strip(): p for p in doc.paragraphs if p.text.strip()}

    for title in (MAIN_TITLE, RESPONSE_TITLE):
        paragraph = paragraphs.get(title)
        if paragraph is None:
            errors.append(f"Falta el encabezado: {title}")
        elif paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            errors.append(f"El encabezado no está centrado: {title}")

    rfi_paragraphs = [
        p for p in doc.paragraphs
        if re.fullmatch(r"RFI-[A-Z]+-\d{4}-\d{3} · REVISIÓN R\d{2}", p.text.strip())
    ]
    if len(rfi_paragraphs) != 1:
        errors.append("Debe existir un único código RFI con revisión.")
    elif rfi_paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.CENTER:
        errors.append("El código y revisión del RFI no están centrados.")

    all_text = "\n".join(p.text for p in doc.paragraphs)
    all_text += "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    for marker in ("**", "$"):
        if marker in all_text:
            errors.append(f"Permanece sintaxis Markdown visible: {marker}")

    for index, table in enumerate(doc.tables, start=1):
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        if layout is None or layout.get(qn("w:type")) != "fixed":
            errors.append(f"La tabla {index} no tiene geometría fija.")

    page_breaks = doc.element.xpath(".//w:br[@w:type='page']")
    if len(page_breaks) != 1:
        errors.append(
            f"Se esperaba un salto explícito entre solicitud y respuesta; encontrados: {len(page_breaks)}."
        )

    if errors:
        raise SystemExit("VALIDACIÓN FALLIDA\n- " + "\n- ".join(errors))

    print(f"VALIDACIÓN ESTRUCTURAL OK: {Path(docx_path).resolve()}")
    print("Encabezados centrados, tablas fijas y un salto solicitud/respuesta.")


def main(argv=None):
    parser = argparse.ArgumentParser(description="Verifica la estructura de un RFI DOCX.")
    parser.add_argument("docx", help="RFI DOCX que se verificará.")
    args = parser.parse_args(argv)
    validate(args.docx)


if __name__ == "__main__":
    main()
