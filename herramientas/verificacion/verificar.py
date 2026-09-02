import sys, io
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt, Cm

PROJECT_ROOT = Path(__file__).resolve().parents[2]
doc = Document(PROJECT_ROOT / "entregables" / "04_borradores" / "IES-001-EST-PRDV-GRA-2026_R01-BORRADOR.docx")

s = doc.sections[0]
print(f"Papel: {s.page_width/360000:.1f} x {s.page_height/360000:.1f} cm")
print(f"Márgenes: T={s.top_margin/360000:.2f} B={s.bottom_margin/360000:.2f} L={s.left_margin/360000:.2f} R={s.right_margin/360000:.2f} cm")
print()

for sn in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
    st = doc.styles[sn]
    sz = st.font.size
    sz_pt = sz / 12700 if sz else 0
    clr = st.font.color.rgb if st.font.color else None
    ls = st.paragraph_format.line_spacing
    sa = st.paragraph_format.space_after
    sa_pt = sa / 12700 if sa else 0
    sb = st.paragraph_format.space_before
    sb_pt = sb / 12700 if sb else 0
    al = st.paragraph_format.alignment
    al_txt = {0: "LEFT", 1: "CENTER", 2: "RIGHT", 3: "JUSTIFY"}.get(al, str(al))
    fn = st.font.name or "inherit"
    print(f"{sn}: {fn} {sz_pt:.0f}pt bold={st.font.bold} color={clr}")
    print(f"  LS={ls} SB={sb_pt:.0f} SA={sa_pt:.0f} align={al_txt}")

print()
for sec in doc.sections:
    for hp in sec.header.paragraphs:
        if hp.text.strip():
            print(f"Header: \"{hp.text[:80]}\"")
print()

# Count paragraphs and tables
n_par = len(doc.paragraphs)
n_tbl = len(doc.tables)
print(f"Párrafos: {n_par}")
print(f"Tablas: {n_tbl}")

# Check table formatting
for ti, t in enumerate(doc.tables[:2]):
    hcell = t.rows[0].cells[0]
    shd = hcell._tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    shd_el = shd.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd') if shd is not None else None
    fill = shd_el.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill') if shd_el is not None else "none"
    vAlign = hcell._tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    vAl_el = vAlign.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}vAlign') if vAlign is not None else None
    va = vAl_el.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if vAl_el is not None else "none"
    print(f"T{ti}: header_bg={fill} vAlign={va}")

# Check page break before H1
for i, p in enumerate(doc.paragraphs):
    if p.style.name == "Heading 1":
        pPr = p._element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
        pb = pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore') if pPr is not None else None
        pb_txt = "YES" if pb is not None else "no"
        print(f"  H1[{i}]: \"{p.text[:40]}\" pageBreakBefore={pb_txt}")

# Check first few paragraphs for alignment
for i, p in enumerate(doc.paragraphs[:20]):
    if p.text.strip():
        al = p.alignment
        al_txt = {0: "LEFT", 1: "CENTER", 2: "RIGHT", 3: "JUSTIFY", None: "inherit"}.get(al, str(al))
        print(f"P[{i}]: align={al_txt} text=\"{p.text[:60]}\"")
