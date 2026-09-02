import sys, io
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
doc = Document(PROJECT_ROOT / "entregables" / "04_borradores" / "CALC-EST-2026-001-R01.docx")

s = doc.sections[0]
print(f"Papel: {s.page_width/360000:.1f}x{s.page_height/360000:.1f} cm")
print(f"Márgenes: T={s.top_margin/360000:.2f} B={s.bottom_margin/360000:.2f} L={s.left_margin/360000:.2f} R={s.right_margin/360000:.2f}")
print()

for sn in ["Normal","Heading 1","Heading 2","Heading 3","Heading 4"]:
    try:
        st=doc.styles[sn]
        sz=st.font.size
        clr=st.font.color.rgb if st.font.color else None
        ls=st.paragraph_format.line_spacing
        print(f"{sn}: {st.font.name} {sz/12700 if sz else 0:.0f}pt bold={st.font.bold} color={clr} LS={ls} SP={st.paragraph_format.space_after/12700 if st.paragraph_format.space_after else 0:.0f}pt")
    except:
        print(f"{sn}: NOT FOUND")

print(f"\nPárrafos: {len(doc.paragraphs)}")
print(f"Tablas: {len(doc.tables)}")
for ti,t in enumerate(doc.tables):
    print(f"  T{ti}: {len(t.rows)}x{len(t.columns)} — {[c.text[:20] for c in t.rows[0].cells]}")

# Check images
n_img = 0
for p in doc.paragraphs:
    for r in p.runs:
        if r._element.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing'):
            n_img += 1
print(f"\nImágenes incrustadas: {n_img}")
