import sys, io
from pathlib import Path
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[2]
doc = Document(PROJECT_ROOT / "entregables" / "04_borradores" / "RFI-EST-2026-001-R00.docx")

s = doc.sections[0]
print(f"Papel: {s.page_width/360000:.1f} x {s.page_height/360000:.1f} cm")
print(f"Márgenes: T={s.top_margin/360000:.2f} B={s.bottom_margin/360000:.2f} L={s.left_margin/360000:.2f} R={s.right_margin/360000:.2f} cm")
print()

for sn in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
    st = doc.styles[sn]
    sz = st.font.size
    sz_pt = sz / 12700 if sz else 0
    print(f"{sn}: {st.font.name} {sz_pt:.0f}pt bold={st.font.bold} color={st.font.color.rgb if st.font.color else None}")
    print(f"  LS={st.paragraph_format.line_spacing} SA={st.paragraph_format.space_after/12700 if st.paragraph_format.space_after else 0:.0f}pt")

print()
print(f"Párrafos: {len(doc.paragraphs)}")
print(f"Tablas: {len(doc.tables)}")
for ti, t in enumerate(doc.tables):
    print(f"  T{ti}: {len(t.rows)}x{len(t.columns)} — headers={[c.text[:25] for c in t.rows[0].cells]}")

print()
# Check paragraph styles
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P[{i}] style={p.style.name} align={p.alignment} \"{p.text[:60]}\"")
